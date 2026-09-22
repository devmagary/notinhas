"""
Sigeduc Auto — Módulo Core
===========================
Funções utilitárias compartilhadas entre a GUI e o CLI.
Inclui: leitura de arquivos, comparação de nomes, controle do Playwright,
        logging, configuração persistente e constantes centralizadas.
"""

import sys
import os
import re
import time
import json
import logging
import unicodedata
import datetime
from pathlib import Path

# =============================================
# Instalação automática de dependências
# =============================================
try:
    from playwright.sync_api import sync_playwright, TimeoutError as PWTimeout
except ImportError:
    print("Playwright não encontrado. Instalando...")
    os.system(f"{sys.executable} -m pip install playwright")
    os.system(f"{sys.executable} -m playwright install chromium")
    from playwright.sync_api import sync_playwright, TimeoutError as PWTimeout

try:
    import openpyxl
except ImportError:
    print("openpyxl não encontrado. Instalando...")
    os.system(f"{sys.executable} -m pip install openpyxl")
    import openpyxl


# =============================================
# Constantes Centralizadas (#19)
# =============================================
URL_FREQUENCIA = "https://sigeduc.educacao.ba.gov.br/sigeduc/portais/docente/docente.jsf"
PASTA_PROJETO = Path(__file__).parent
PERFIL_NAVEGADOR = PASTA_PROJETO / "perfil_navegador"
CONFIG_PATH = PASTA_PROJETO / "config.json"
LOG_PATH = PASTA_PROJETO / "sigeduc_auto.log"

TIMEOUT_TABELA = 10000       # ms — espera carregar tabela de alunos
TIMEOUT_PAGINA = 60000       # ms — espera carregar página completa
TIMEOUT_NETWORKIDLE = 15000  # ms — espera networkidle pós-gravação
TIMEOUT_CLICK = 1500         # ms — espera por clique em checkbox/input
ESPERA_CARREGAMENTO = 3      # s  — sleep após carregamento de página
ESPERA_POS_GRAVACAO = 2      # s  — sleep pós-gravação
ESPERA_PRE_CHECKBOX = 1      # s  — sleep antes de manipular checkboxes
SLOW_MO = 200                # ms — delay do Playwright entre ações
MAX_TENTATIVAS_CONTEUDO = 5  # tentativas para obter conteúdo da página
MAX_TENTATIVAS_DATA = 3      # tentativas para encontrar data no calendário

MESES = [
    "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
    "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"
]


# =============================================
# Logging Persistente (#20)
# =============================================
logger = logging.getLogger("sigeduc_auto")
logger.setLevel(logging.DEBUG)

# Evita adicionar handlers duplicados se o módulo for recarregado
if not logger.handlers:
    # Handler: arquivo de log
    fh = logging.FileHandler(str(LOG_PATH), encoding="utf-8")
    fh.setLevel(logging.DEBUG)
    fh.setFormatter(logging.Formatter("%(asctime)s | %(levelname)-7s | %(message)s", datefmt="%Y-%m-%d %H:%M:%S"))
    logger.addHandler(fh)

    # Handler: console (fallback, se não estiver usando a GUI)
    ch = logging.StreamHandler(sys.stdout)
    ch.setLevel(logging.INFO)
    ch.setFormatter(logging.Formatter("%(message)s"))
    logger.addHandler(ch)


# =============================================
# Configuração Persistente (#11)
# =============================================
def salvar_config(dados: dict):
    """Salva configurações do usuário em um arquivo JSON local."""
    try:
        with open(CONFIG_PATH, "w", encoding="utf-8") as f:
            json.dump(dados, f, ensure_ascii=False, indent=2)
    except Exception:
        logger.warning("Não foi possível salvar configurações.")


def carregar_config() -> dict:
    """Carrega configurações salvas do arquivo JSON."""
    try:
        if CONFIG_PATH.exists():
            with open(CONFIG_PATH, "r", encoding="utf-8") as f:
                return json.load(f)
    except Exception:
        logger.warning("Não foi possível carregar configurações anteriores.")
    return {}


# =============================================
# Comparação Robusta de Nomes e Desambiguação (#1 e #15)
# =============================================
PREPOSICOES_NOME = {"DE", "DA", "DO", "DAS", "DOS", "E"}
SUFIXOS_NOME = {"JUNIOR", "JR", "FILHO", "NETO", "SOBRINHO", "SEGUNDO", "TERCEIRO"}


def normalizar_nome(nome: str) -> str:
    """Remove acentos, caracteres especiais e normaliza espaços."""
    if not nome:
        return ""
    nome = unicodedata.normalize("NFKD", str(nome))
    nome = "".join(c for c in nome if not unicodedata.combining(c))
    nome = re.sub(r"[^A-Za-z\s]", " ", nome)
    return " ".join(nome.upper().split())


def canonicalizar_palavra(p: str) -> str:
    """
    Normaliza variações ortográficas e fonéticas frequentes em nomes brasileiros
    (ex: LUIS/LUIZ, SOUZA/SOUSA, MATHEUS/MATEUS, HELOYSA/ELOISA, VICTORIA/VITORIA).
    """
    p = p.upper()
    if p.startswith("H") and len(p) > 1 and p[1] in "AEIOU":
        p = p[1:]
    p = p.replace("TH", "T").replace("PH", "F").replace("CH", "X")
    p = p.replace("CT", "T").replace("PT", "T").replace("SC", "S")
    p = p.replace("Y", "I").replace("W", "V")
    p = p.replace("Z", "S")
    p = p.replace("K", "C")
    p = re.sub(r"([B-DF-HJ-NP-TV-Z])\1+", r"\1", p)
    return p


def extrair_palavras_relevantes(nome: str, canon: bool = True) -> list[str]:
    """Retorna a lista de palavras substantivas do nome (excluindo preposições)."""
    palavras = [p for p in normalizar_nome(nome).split() if p not in PREPOSICOES_NOME]
    if canon:
        return [canonicalizar_palavra(p) for p in palavras]
    return palavras


def nomes_sao_iguais(nome_a: str, nome_b: str) -> bool:
    """
    Verifica se dois nomes representam com certeza absoluta a mesma pessoa,
    eliminando rigorosamente falsos positivos entre estudantes com prenomes iguais
    (ex: MARIA VITORIA SANTOS LISBOA vs MARIA VITORIA TOSTA DO AMARAL -> False).
    """
    a_norm = normalizar_nome(nome_a)
    b_norm = normalizar_nome(nome_b)
    if not a_norm or not b_norm:
        return False
    if a_norm == b_norm:
        return True

    p_a = extrair_palavras_relevantes(nome_a, canon=True)
    p_b = extrair_palavras_relevantes(nome_b, canon=True)
    if not p_a or not p_b:
        return False

    if p_a == p_b:
        return True

    # Primeiro nome DEVE ser idêntico
    if p_a[0] != p_b[0]:
        return False

    # Trata sufixos de parentesco (Junior, Filho, Neto, etc.)
    ult_a = p_a[-1]
    ult_b = p_b[-1]
    if ult_a in SUFIXOS_NOME and ult_b not in SUFIXOS_NOME:
        p_a = p_a[:-1]
        if not p_a:
            return False
        ult_a = p_a[-1]
    elif ult_b in SUFIXOS_NOME and ult_a not in SUFIXOS_NOME:
        p_b = p_b[:-1]
        if not p_b:
            return False
        ult_b = p_b[-1]

    # Último sobrenome de família DEVE ser estritamente compatível
    if len(p_a) >= 2 and len(p_b) >= 2:
        if ult_a != ult_b:
            eh_abrev = (len(ult_a) == 1 and ult_b.startswith(ult_a)) or (len(ult_b) == 1 and ult_a.startswith(ult_b))
            if not eh_abrev:
                return False

    # Verifica conflito de palavras substantivas intermediárias
    set_a = set(p_a)
    set_b = set(p_b)
    so_a = set_a - set_b
    so_b = set_b - set_a

    conflitos_a = {pa for pa in so_a if not (len(pa) == 1 and any(pb.startswith(pa) for pb in so_b))}
    conflitos_b = {pb for pb in so_b if not (len(pb) == 1 and any(pa.startswith(pb) for pa in so_a))}

    # Se ambos possuem palavras substantivas completas conflitantes (ex: ARAUJO vs OLIVEIRA),
    # então NÃO é a mesma pessoa!
    if conflitos_a and conflitos_b:
        return False

    menor = p_a if len(p_a) <= len(p_b) else p_b
    maior = p_b if len(p_a) <= len(p_b) else p_a
    if len(menor) < 2:
        return False

    palavras_comuns = set(menor) & set(maior)
    if len(palavras_comuns) / len(maior) >= 0.65:
        return True

    return False


def nomes_correspondem(nome_tela: str, nome_lista: str) -> bool:
    """
    Compara dois nomes garantindo que nomes parecidos de estudantes diferentes
    nunca gerem correspondência indevida.
    """
    a_norm = normalizar_nome(nome_tela)
    b_norm = normalizar_nome(nome_lista)
    if not a_norm or not b_norm:
        return False
    if a_norm == b_norm:
        return True

    # Se um nome está contido como substring de palavras completas no outro
    padrao_b = r"\b" + re.escape(b_norm) + r"\b"
    if re.search(padrao_b, a_norm):
        return True
    padrao_a = r"\b" + re.escape(a_norm) + r"\b"
    if re.search(padrao_a, b_norm):
        return True

    return nomes_sao_iguais(nome_tela, nome_lista)


def encontrar_melhor_aluno(texto_tela: str, candidatos: list[str]) -> str | None:
    """
    Localiza o aluno correto da lista de candidatos a partir do texto/linha da tela.
    Avalia TODOS os candidatos e seleciona o match de maior certeza,
    evitando que o primeiro aluno que passe no filtro sobrescreva alunos com nomes parecidos.
    """
    if not texto_tela or not candidatos:
        return None

    texto_norm = normalizar_nome(texto_tela)

    # 1. Match exato direto
    for cand in candidatos:
        cand_norm = normalizar_nome(cand)
        if cand_norm == texto_norm:
            return cand

    # 2. O nome do candidato está contido integralmente como sequência de palavras na linha
    matches_contidos = []
    for cand in candidatos:
        cand_norm = normalizar_nome(cand)
        padrao = r"\b" + re.escape(cand_norm) + r"\b"
        if re.search(padrao, texto_norm):
            matches_contidos.append((len(cand_norm), cand))

    if matches_contidos:
        matches_contidos.sort(reverse=True, key=lambda x: x[0])
        return matches_contidos[0][1]

    # 3. Match canônico (nomes_sao_iguais)
    matches_canon = []
    for cand in candidatos:
        if nomes_sao_iguais(cand, texto_tela):
            matches_canon.append(cand)

    if len(matches_canon) == 1:
        return matches_canon[0]
    elif len(matches_canon) > 1:
        # Em caso de múltiplos matches canônicos, seleciona o de maior similaridade de palavras
        palavras_tela = set(extrair_palavras_relevantes(texto_tela, canon=True))
        melhor = None
        melhor_score = -1
        for cand in matches_canon:
            palavras_cand = set(extrair_palavras_relevantes(cand, canon=True))
            score = len(palavras_tela & palavras_cand)
            if score > melhor_score:
                melhor_score = score
                melhor = cand
        return melhor

    return None


def extrair_nome_da_linha(linha) -> str:
    """
    Extrai o nome do aluno da linha da tabela do SIGEduc procurando
    especificamente pelo elemento link <a> ou pela célula <td> de nome.
    """
    try:
        # 1. Procura por links <a> (no SIGEduc, o nome do estudante é link clicável)
        links = linha.query_selector_all("a")
        for link in links:
            txt = link.inner_text().strip()
            partes = txt.split()
            if len(partes) >= 2 and len(txt) >= 5:
                txt_upper = txt.upper()
                termos_proibidos = [
                    "VOLTAR", "EDITAR", "EXCLUIR", "VISUALIZAR", "ALTERAR", "GRAVAR", "SALVAR", 
                    "HISTORICO", "FREQUENCIA", "DETALHES", "IMPRIMIR", "ANEXO", "AJUDA"
                ]
                if not any(t in txt_upper for t in termos_proibidos):
                    return txt

        # 2. Procura por células <td> específicas sem inputs
        tds = linha.query_selector_all("td")
        for td in tds:
            if td.query_selector("input, select, button, textarea"):
                continue
            txt = td.inner_text().strip()
            partes = txt.split()
            if len(partes) >= 2 and len(txt) >= 5:
                if partes[0].isalpha() and partes[1].isalpha():
                    txt_upper = txt.upper()
                    if not any(kw in txt_upper for kw in ["MATRÍCULA", "MATRICULA", "SITUAÇÃO", "SITUACAO"]):
                        return txt
    except Exception:
        pass

    try:
        return linha.inner_text().strip()
    except Exception:
        return ""


def encontrar_na_lista(nome_tela: str, nomes: list[str]) -> bool:
    """Verifica se o nome da tela corresponde a algum aluno da lista."""
    return encontrar_melhor_aluno(nome_tela, nomes) is not None


# =============================================
# Leitura de Arquivos
# =============================================
def ler_nomes(caminho_arquivo: str) -> list[str]:
    """Lê nomes de um arquivo CSV/TXT ou XLSX (coluna A). (#8)"""
    if caminho_arquivo.lower().endswith(".xlsx"):
        return _ler_nomes_xlsx(caminho_arquivo)
    else:
        return _ler_nomes_texto(caminho_arquivo)


def _ler_nomes_texto(caminho_arquivo: str) -> list[str]:
    """Lê nomes do arquivo CSV/TXT (detecta automaticamente se a primeira coluna é número de chamada)."""
    nomes = []

    # Tenta ler com diferentes codificações (comum CSV do Excel vir como latin-1)
    encodings = ["utf-8-sig", "utf-8", "latin-1", "cp1252"]
    conteudo = ""
    for enc in encodings:
        try:
            with open(caminho_arquivo, "r", encoding=enc) as f:
                conteudo = f.read()
            break
        except UnicodeDecodeError:
            continue

    for linha in conteudo.splitlines():
        linha = linha.strip()
        if not linha:
            continue
        # Divide por vírgula, ponto-e-vírgula ou tab
        partes = [p.strip().upper() for p in re.split(r"[,;|\t]", linha) if p.strip()]
        if partes:
            # Se a 1ª coluna for só número de chamada (ex: "1", "02"), pega a 2ª coluna
            if partes[0].isdigit() and len(partes) > 1:
                nome = partes[1]
            else:
                nome = partes[0]
            if len(nome) > 1 and any(c.isalpha() for c in nome):
                nomes.append(nome)

    return nomes


def _ler_nomes_xlsx(caminho_arquivo: str) -> list[str]:
    """Lê nomes da coluna A (ou B se a coluna A for apenas número de chamada) de um arquivo XLSX."""
    nomes = []
    wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
    try:
        planilha = wb.active
        col_idx = 0
        
        # Inspeciona as primeiras linhas para verificar se a Coluna A é número de chamada (1, 2, 3...)
        linhas_amostra = list(planilha.iter_rows(min_row=1, max_row=10, max_col=3, values_only=True))
        cont_nomes_col0 = 0
        cont_nomes_col1 = 0
        for l in linhas_amostra[1:]:  # pula cabeçalho
            if l:
                v0 = str(l[0]).strip() if len(l) > 0 and l[0] is not None else ""
                v1 = str(l[1]).strip() if len(l) > 1 and l[1] is not None else ""
                if any(c.isalpha() for c in v0) and len(v0.split()) >= 2:
                    cont_nomes_col0 += 1
                if any(c.isalpha() for c in v1) and len(v1.split()) >= 2:
                    cont_nomes_col1 += 1

        if cont_nomes_col1 > cont_nomes_col0:
            col_idx = 1
            logger.info("  ✓ Nomes dos alunos identificados na Coluna B do Excel.")
        else:
            col_idx = 0

        for linha in planilha.iter_rows(min_row=2, max_col=col_idx + 1, values_only=True):
            if linha and len(linha) > col_idx and linha[col_idx]:
                nome = str(linha[col_idx]).strip().upper()
                if len(nome) > 1 and any(c.isalpha() for c in nome):
                    nomes.append(nome)
    finally:
        wb.close()
    return nomes


def letra_para_indice(letra: str) -> int:
    """Converte letras de colunas do Excel (A, B, Z, AA) em índices (A=0, B=1)."""
    letra = letra.strip().upper()
    soma = 0
    for i, c in enumerate(reversed(letra)):
        if "A" <= c <= "Z":
            soma += (ord(c) - 64) * (26 ** i)
    return max(0, soma - 1)


VALORES_VAZIOS = {
    "", "-", "--", "---", "n/a", "na", "null", "none",
    "sem nota", "s/n", "sn", "sem avaliacao", "sem avaliação",
    "pendente", "fv", "falta", "faltou", "ausente", "disp", "dispensado"
}


def formatar_nota(val, formato: str = "ponto") -> str:
    """
    Formata o valor da nota para o padrão aceito pelo Sigeduc.
    - 'ponto': padrão oficial do Sigeduc (ex: '3.5', '3.0', '10.0').
    - 'virgula': formatação com vírgula decimal (ex: '3,5', '3,0', '10,0').
    - 'original': mantém o texto original sem alteração de pontuação (se numérico).
    Valores textuais não numéricos ('Sem nota', 'FV', '-', etc.) são convertidos para ''
    para não injetar texto em campos numéricos do portal.
    """
    if val is None:
        return ""
    texto = str(val).strip()
    texto_lower = texto.lower()

    # Sentinelas e valores textuais comuns em planilhas que indicam ausência de nota
    if texto_lower in VALORES_VAZIOS:
        return ""

    # Normaliza vírgula para ponto a fim de fazer o parse numérico
    texto_num = texto.replace(",", ".")
    try:
        num = float(texto_num)
        if formato == "original":
            return texto
        # Formata com 1 casa decimal (padrão das máscaras do Sigeduc: 3.0, 7.5, 10.0)
        formatado = f"{num:.1f}"
        if formato == "virgula":
            return formatado.replace(".", ",")
        return formatado
    except ValueError:
        # Se for qualquer outro texto não numérico, trata como sem nota para não quebrar o site
        logger.info(f"  ⚠ Valor não numérico ignorado (deixado em branco): '{texto}'")
        return ""



def eh_falta_vinculada(val) -> bool:
    """
    Retorna True se a célula representa ausência de nota (Falta Vinculada).
    No SIGEduc, 'Sem nota', 'FV', campos vazios ou não numéricos são considerados FV.
    """
    if val is None:
        return True
    txt = str(val).strip().lower()
    if txt in VALORES_VAZIOS or txt == "":
        return True
    try:
        float(txt.replace(",", "."))
        return False
    except ValueError:
        return True


class DadosAluno(dict):
    """
    Objeto com compatibilidade total com listas:
    Permite:
      - dados_notas[aluno][i] -> nota formatada (ex: '2.9' ou '')
      - dados_notas[aluno].fv -> lista de booleanos indicando Falta Vinculada por coluna
      - dados_notas[aluno].tem_fv -> True/False
      - len(dados_notas[aluno]) -> quantidade de notas
      - iter(dados_notas[aluno]) -> iterador das notas
    """
    def __init__(self, notas: list[str], fv: list[bool]):
        super().__init__(notas=notas, fv=fv)
        self.notas = notas
        self.fv = fv
        self.tem_fv = any(fv)

    def __getitem__(self, item):
        if isinstance(item, int) or isinstance(item, slice):
            return self.notas[item]
        return super().__getitem__(item)

    def __len__(self):
        return len(self.notas)

    def __iter__(self):
        return iter(self.notas)


def ler_notas_xlsx(caminho_arquivo: str, letras_colunas: list[str], formato: str = "ponto") -> dict:
    """Lê as notas do arquivo XLSX e retorna { 'NOME': DadosAluno(notas=[...], fv=[...]) }. (#16 wb.close)"""
    notas = {}
    wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
    try:
        planilha = wb.active
        indices_colunas = [letra_para_indice(l) for l in letras_colunas]

        for linha in planilha.iter_rows(min_row=2, values_only=True):
            if not linha or not linha[0]:
                continue
            nome = str(linha[0]).strip().upper()

            valores = []
            flags_fv = []
            for idx in indices_colunas:
                val = linha[idx] if len(linha) > idx else None
                flags_fv.append(eh_falta_vinculada(val))
                valores.append(formatar_nota(val, formato=formato))

            notas[nome] = DadosAluno(notas=valores, fv=flags_fv)
    finally:
        wb.close()
    return notas



# =============================================
# Parsing e Validação de Datas (#7)
# =============================================
def parsear_datas(texto_datas: str) -> list[tuple[int, int]]:
    """
    Converte texto com datas (ex: '11/03, 18/03; 01/04') em [(11, 3), (18, 3), (1, 4)].
    Suporta múltiplos separadores (vírgula, ponto-e-vírgula, quebras de linha e espaços),
    valida se o dia/mês existem no calendário e ordena cronologicamente sem repetição.
    """
    datas_encontradas = []
    ano_atual = datetime.date.today().year

    # Divide por vírgula, ponto e vírgula, quebras de linha ou múltiplos espaços
    partes = [p.strip() for p in re.split(r"[,;\s\n\r]+", texto_datas.strip()) if p.strip()]

    for parte in partes:
        if "/" not in parte:
            continue
        pedacos = parte.split("/")
        if len(pedacos) < 2:
            continue
        try:
            d, m = int(pedacos[0]), int(pedacos[1])
            # Valida se a data é real no calendário gregoriano
            datetime.date(ano_atual, m, d)
            if (d, m) not in datas_encontradas:
                datas_encontradas.append((d, m))
        except (ValueError, IndexError):
            logger.warning(f"  ⚠ Data inválida ignorada: {parte}")

    # Ordena cronologicamente por mês e dia
    datas_ordenadas = sorted(datas_encontradas, key=lambda x: (x[1], x[0]))
    return datas_ordenadas



# =============================================
# Controle do Navegador (Playwright)
# =============================================
def obter_conteudo_seguro(page) -> str:
    """Obtém conteúdo da página com backoff exponencial. (#14)"""
    for i in range(MAX_TENTATIVAS_CONTEUDO):
        try:
            return page.content()
        except Exception:
            time.sleep(min(2 ** i, 8))
    return ""


def trazer_navegador_para_frente():
    """Traz o navegador Chrome do Sigeduc para o primeiro plano no Windows. (#17 ctypes fix)"""
    try:
        import ctypes
        EnumWindows = ctypes.windll.user32.EnumWindows
        # Corrigido: c_void_p em vez de c_pointer (que não existe)
        EnumWindowsProc = ctypes.WINFUNCTYPE(ctypes.c_bool, ctypes.c_void_p, ctypes.c_void_p)
        GetWindowText = ctypes.windll.user32.GetWindowTextW
        GetWindowTextLength = ctypes.windll.user32.GetWindowTextLengthW
        IsWindowVisible = ctypes.windll.user32.IsWindowVisible
        SetForegroundWindow = ctypes.windll.user32.SetForegroundWindow
        ShowWindow = ctypes.windll.user32.ShowWindow

        def foreach_window(hwnd, lParam):
            if IsWindowVisible(hwnd):
                length = GetWindowTextLength(hwnd)
                if length > 0:
                    buff = ctypes.create_unicode_buffer(length + 1)
                    GetWindowText(hwnd, buff, length + 1)
                    title = buff.value.lower()
                    if any(term in title for term in ["sigeduc", "docente", "frequência", "diário", "chrome"]):
                        ShowWindow(hwnd, 9)  # SW_RESTORE
                        SetForegroundWindow(hwnd)
                        return False
            return True

        EnumWindows(EnumWindowsProc(foreach_window), 0)
    except Exception:
        pass


def tratar_dialogo_confirmacao(page):
    """Configura handler para aceitar automaticamente os dialogs do site. (#13 truncamento fix)"""
    def handler(dialog):
        msg = dialog.message
        msg_display = f"{msg[:60]}..." if len(msg) > 60 else msg
        logger.info(f"  ✓ Diálogo detectado: '{msg_display}' → Aceitando.")
        dialog.accept()
    page.on("dialog", handler)


def aguardar_pos_gravacao(page):
    """Aguarda o site processar a gravação (redirect ou AJAX)."""
    logger.info("  Aguardando processamento do site...")
    try:
        page.wait_for_load_state("networkidle", timeout=TIMEOUT_NETWORKIDLE)
    except PWTimeout:
        pass
    time.sleep(ESPERA_POS_GRAVACAO)


# =============================================
# Detecção de Calendário
# =============================================
def e_tabela_calendario(tabela) -> bool:
    """Verifica se uma tabela possui características de um calendário mensal."""
    try:
        texto = normalizar_nome(tabela.inner_text() or "")
        dias_semana = ["DOM", "SEG", "TER", "QUA", "QUI", "SEX", "SAB"]
        contador_dias = sum(1 for d in dias_semana if d in texto)
        if contador_dias >= 4:
            links = tabela.query_selector_all("a")
            if len(links) >= 5:
                return True
    except Exception:
        pass
    return False


def garantir_pagina_calendario(page) -> bool:
    """
    Verifica se estamos na página do calendário e navega de volta de forma segura e não destrutiva.
    Evita page.goto(URL_FREQUENCIA) que desloga ou sai da turma ativa.
    """
    try:
        page.wait_for_load_state("networkidle", timeout=5000)
    except PWTimeout:
        pass

    # 1. Verifica se já está em uma página contendo calendário mensal
    for tab in page.query_selector_all("table"):
        if e_tabela_calendario(tab):
            return True

    logger.info("  ⚠ Calendário não visível. Procurando botão de retorno...")

    # 2. Procura botões, links ou inputs de retorno ao diário / frequência
    elementos = page.query_selector_all("input[type='button'], input[type='submit'], button, a")
    for el in elementos:
        try:
            val = normalizar_nome(el.get_attribute("value") or el.inner_text() or "")
        except Exception:
            continue
        if any(term in val for term in ["VOLTAR", "RETORNAR", "CALENDARIO", "DIARIO DE CLASSE"]):
            try:
                el.click()
                time.sleep(ESPERA_CARREGAMENTO)
                try:
                    page.wait_for_load_state("networkidle", timeout=TIMEOUT_TABELA)
                except PWTimeout:
                    pass
                for tab in page.query_selector_all("table"):
                    if e_tabela_calendario(tab):
                        return True
            except Exception:
                pass

    # 3. Tenta voltar no histórico do navegador se o formulário permitir
    try:
        page.go_back(wait_until="networkidle", timeout=5000)
        time.sleep(ESPERA_CARREGAMENTO)
        for tab in page.query_selector_all("table"):
            if e_tabela_calendario(tab):
                return True
    except Exception:
        pass

    return False


def encontrar_dia_no_calendario(page, dia: int, mes: int) -> bool:
    """Encontra e clica no dia correto dentro do mês correto no calendário."""
    nome_mes = MESES[mes - 1]
    nome_mes_norm = normalizar_nome(nome_mes)
    logger.info(f"  Procurando dia {dia:02d} no mês de {nome_mes}...")

    tabelas = page.query_selector_all("table")
    for tabela in tabelas:
        try:
            texto_tabela = normalizar_nome(tabela.inner_text() or "")
        except Exception:
            continue
        if nome_mes_norm not in texto_tabela:
            continue

        links = tabela.query_selector_all("a")
        for link in links:
            try:
                txt_link = link.inner_text().strip()
                # Aceita tanto formato simples '5' quanto formatado com zero à esquerda '05'
                if txt_link in (str(dia), f"{dia:02d}"):
                    logger.info(f"  ✓ Dia {dia:02d} encontrado e clicável. Clicando...")
                    link.click()
                    return True
            except Exception:
                continue

    logger.warning(f"  ✗ Dia {dia:02d} de {nome_mes} não é clicável no calendário.")
    return False



# =============================================
# Marcação de Presenças
# =============================================
def marcar_presencas(page, nomes: list[str], modo: str) -> int:
    """Marca ou desmarca os checkboxes de presença para os alunos da lista."""
    logger.info("  Aguardando tabela de alunos...")
    try:
        page.wait_for_selector("#tableAlunos tbody tr", timeout=TIMEOUT_TABELA)
    except PWTimeout:
        logger.warning("  ⚠ Tabela de alunos não carregou. Tentando continuar...")

    time.sleep(ESPERA_PRE_CHECKBOX)

    linhas = page.query_selector_all("#tableAlunos tbody tr")
    logger.info(f"  Alunos encontrados na tela: {len(linhas)}")

    marcados = 0
    for linha in linhas:
        celula_nome = linha.query_selector("td:nth-child(3)")
        nome_tela = celula_nome.inner_text().strip().upper() if celula_nome else ""
        if not nome_tela:
            nome_tela = extrair_nome_da_linha(linha)
        if not nome_tela:
            continue

        # Comparação robusta de nomes e desambiguação (#1)
        na_lista = encontrar_na_lista(nome_tela, nomes)

        checkboxes = linha.query_selector_all('input[type="checkbox"]')

        # Filtra apenas checkboxes habilitados
        habilitados = []
        tem_desabilitado = False
        for cb in checkboxes:
            try:
                if cb.is_disabled() or cb.get_attribute("disabled") is not None:
                    tem_desabilitado = True
                else:
                    habilitados.append(cb)
            except Exception:
                tem_desabilitado = True

        if tem_desabilitado:
            logger.info(f"  ⚠ {nome_tela}: Registro bloqueado/justificado (não alterado).")

        if len(habilitados) >= 1:
            if modo == "1":
                # Faltosos na lista: marcar falta na lista, presença no resto
                if na_lista:
                    for cb in habilitados:
                        try:
                            if not cb.is_checked(): cb.click(timeout=TIMEOUT_CLICK)
                        except Exception:
                            pass
                    marcados += 1
                else:
                    for cb in habilitados:
                        try:
                            if cb.is_checked(): cb.click(timeout=TIMEOUT_CLICK)
                        except Exception:
                            pass

            elif modo == "2":
                # Presentes na lista: presença na lista, falta no resto
                if na_lista:
                    for cb in habilitados:
                        try:
                            if cb.is_checked(): cb.click(timeout=TIMEOUT_CLICK)
                        except Exception:
                            pass
                else:
                    for cb in habilitados:
                        try:
                            if not cb.is_checked(): cb.click(timeout=TIMEOUT_CLICK)
                        except Exception:
                            pass
                    marcados += 1

            elif modo == "3":
                # Presença na lista, ignora o resto
                if na_lista:
                    for cb in habilitados:
                        try:
                            if cb.is_checked(): cb.click(timeout=TIMEOUT_CLICK)
                        except Exception:
                            pass
                    marcados += 1

            elif modo == "4":
                # Falta em quem NÃO está na lista, ignora a lista
                if not na_lista:
                    for cb in habilitados:
                        try:
                            if not cb.is_checked(): cb.click(timeout=TIMEOUT_CLICK)
                        except Exception:
                            pass
                    marcados += 1

            elif modo == "5":
                # 100% Presentes: Garante presença para todos (desmarca checkboxes de falta habilitados)
                for cb in habilitados:
                    try:
                        if cb.is_checked(): cb.click(timeout=TIMEOUT_CLICK)
                    except Exception:
                        pass
                marcados += 1

    return marcados


# =============================================
# Gerenciamento de Abas de Unidades
# =============================================
def garantir_unidade_ativa(page, unidade: str | int = 1, timeout: float = 8.0) -> bool:
    """
    Garante que a aba correspondente à unidade informada esteja ativa no SIGEduc.
    Funciona tanto na tela de lançamento de notas quanto na tela de Falta Vinculada
    ('VERIFICAR AVALIAÇÕES NÃO REALIZADAS').

    Verifica classes RichFaces/JSF e estilos visuais (texto branco / fundo azul).
    Caso a aba já esteja ativa, não efetua clique desnecessário.
    """
    if not unidade:
        return True

    match = re.search(r'([1-4])', str(unidade))
    if not match:
        logger.warning(f"  ⚠ Identificador de unidade inválido: '{unidade}'. Mantendo aba atual.")
        return False

    num_unidade = int(match.group(1))
    texto_unidade = f"{num_unidade}ª Unidade"

    logger.info(f"  Verificando se a {texto_unidade} está selecionada...")

    js_inspecionar_e_clicar = """
    (targetNum) => {
        const candidatos = Array.from(document.querySelectorAll("a, td, div, span, button, li"));
        const regexUnidade = new RegExp(`^\\\\s*${targetNum}[ªºa]?\\\\s*unidade\\\\b`, 'i');
        const regexGeral = new RegExp(`\\\\b${targetNum}[ªºa]?\\\\s*unidade\\\\b`, 'i');

        let tabElement = null;
        for (const el of candidatos) {
            if (el.children.length > 2) continue;
            const txt = (el.innerText || el.textContent || "").trim();
            if (regexUnidade.test(txt) || regexGeral.test(txt)) {
                tabElement = el;
                break;
            }
        }

        if (!tabElement) {
            const idMap = { 1: "tab1", 2: "tab1j_id_1", 3: "tab1j_id_2", 4: "tab1j_id_3" };
            const alvoId = idMap[targetNum];
            if (alvoId) {
                tabElement = document.querySelector(`a[href*='${alvoId}'], [id*='${alvoId}']`);
            }
        }

        if (!tabElement) {
            return { sucesso: false, motivo: "Elemento da aba não localizado na página" };
        }

        function verificarAtivo(el) {
            let cur = el;
            for (let i = 0; i < 4 && cur; i++) {
                const cls = (cur.className || "").toString().toLowerCase();
                if (cls.includes("rich-tab-active") || cls.includes("dr-tbpnl-tb-act") || 
                    cls.includes("rich-tabhdr-cell-active") || cls.includes("tab-active") || 
                    cls.includes("active") || cur.getAttribute("aria-selected") === "true") {
                    return true;
                }
                const st = window.getComputedStyle(cur);
                if (st.color === "rgb(255, 255, 255)" || st.color === "#ffffff") {
                    return true;
                }
                const bg = st.backgroundColor;
                const m = bg.match(/\\d+/g);
                if (m && m.length >= 3) {
                    const [r, g, b] = m.map(Number);
                    if (b > 100 && b > r && b > g) {
                        return true;
                    }
                }
                cur = cur.parentElement;
            }
            return false;
        }

        if (verificarAtivo(tabElement)) {
            return { sucesso: true, ja_ativa: true };
        }

        const clicavel = tabElement.tagName === "A" || tabElement.tagName === "BUTTON" 
            ? tabElement 
            : (tabElement.querySelector("a, button") || tabElement);

        clicavel.scrollIntoView({ block: "center", inline: "center" });
        clicavel.click();

        return { sucesso: true, ja_ativa: false };
    }
    """

    try:
        resultado = page.evaluate(js_inspecionar_e_clicar, num_unidade)
    except Exception as e:
        logger.warning(f"  ⚠ Erro ao inspecionar abas de unidade: {e}")
        return False

    if not resultado.get("sucesso"):
        logger.info(f"  Aba da {texto_unidade} não localizada na tela (pode não haver abas nesta visualização).")
        return False

    if resultado.get("ja_ativa"):
        logger.info(f"  ✓ {texto_unidade} já está ativa e pronta!")
        return True

    logger.info(f"  ✓ Aba {texto_unidade} clicada. Aguardando recarregamento dos dados...")

    try:
        page.wait_for_load_state("networkidle", timeout=5000)
    except Exception:
        pass
    time.sleep(1.5)

    try:
        revalida = page.evaluate(js_inspecionar_e_clicar, num_unidade)
        if revalida.get("ja_ativa"):
            logger.info(f"  ✓ {texto_unidade} ativada com sucesso!")
            return True
    except Exception:
        pass

    return True


# =============================================
# Lançamento de Notas
# =============================================
def lancar_notas(page, dados_notas: dict, unidade: str | int = None) -> int:
    """Preenche as notas na página para os alunos fornecidos no dicionário."""
    if unidade:
        garantir_unidade_ativa(page, unidade)

    logger.info("  Aguardando tabela de notas...")
    try:
        page.wait_for_selector(
            "table tbody tr input:not([type='hidden']):not([type='checkbox']):not([type='radio'])",
            timeout=TIMEOUT_TABELA,
        )
    except Exception:
        logger.warning("  ⚠ Tabela de notas não carregou. Tentando continuar...")

    time.sleep(ESPERA_PRE_CHECKBOX)

    linhas = page.query_selector_all("table tbody tr")
    logger.info("  Verificando linhas na tabela para notas...")

    alunos_preenchidos = 0
    alunos_ja_processados = set()
    candidatos_nomes = list(dados_notas.keys())

    for linha in linhas:
        try:
            texto_linha = linha.inner_text().strip().upper()
        except Exception:
            continue

        if not texto_linha:
            continue

        # Localização robusta do aluno pelo nome do link/célula ou linha inteira
        nome_tela = extrair_nome_da_linha(linha)
        aluno_encontrado = encontrar_melhor_aluno(nome_tela, candidatos_nomes)
        if not aluno_encontrado and texto_linha and texto_linha != nome_tela:
            aluno_encontrado = encontrar_melhor_aluno(texto_linha, candidatos_nomes)

        if not aluno_encontrado:
            continue

        # Evita associar o mesmo aluno a mais de uma linha da tabela
        if aluno_encontrado in alunos_ja_processados:
            logger.warning(f"  ⚠ Aluno '{aluno_encontrado}' já foi associado a outra linha. Pulando duplicidade.")
            continue

        alunos_ja_processados.add(aluno_encontrado)

        notas = dados_notas[aluno_encontrado]
        qtd_notas = len(notas)

        # Prioriza campos de avaliação da tela do Sigeduc (ex: class="avaliacao"),
        # senão busca inputs visíveis ignorando campos calculados (ex: .noEdit)
        inputs_avaliacao = [inp for inp in linha.query_selector_all("input.avaliacao") if inp.is_visible()]
        if inputs_avaliacao and len(inputs_avaliacao) >= qtd_notas:
            inputs = inputs_avaliacao
        else:
            inputs_brutos = linha.query_selector_all(
                "input:not([type='hidden']):not([type='checkbox']):not([type='radio']):not(.noEdit)"
            )
            inputs = [inp for inp in inputs_brutos if inp.is_visible()]

        if len(inputs) >= qtd_notas:
            detalhes = []
            for i in range(qtd_notas):
                val = notas[i]
                if val != "":
                    try:
                        if not inputs[i].is_disabled() and inputs[i].get_attribute("disabled") is None:
                            inputs[i].fill(val, timeout=TIMEOUT_CLICK)
                            detalhes.append(f"Col {i+1}: {val}")
                        else:
                            detalhes.append(f"Col {i+1}: bloqueada")
                    except Exception as e:
                        logger.warning(f"  ⚠ Erro ao preencher nota de {aluno_encontrado} (coluna {i+1}): {e}")
                else:
                    detalhes.append(f"Col {i+1}: em branco")

            logger.info(f"  ✓ {aluno_encontrado} -> {', '.join(detalhes)}")
            alunos_preenchidos += 1


    return alunos_preenchidos


# =============================================
# Gravação de Dados
# =============================================
def _clicar_primeiro_visivel(page, seletores_ou_textos: list) -> bool:
    """Busca e clica no primeiro elemento visível que corresponda à lista de seletores ou textos."""
    for item in seletores_ou_textos:
        if any(c in item for c in ["[", ":", ".", "#"]):
            try:
                for el in page.query_selector_all(item):
                    if el.is_visible():
                        logger.info(f"  ✓ Elemento clicado via seletor: {item}")
                        el.click()
                        return True
            except Exception:
                pass
        seletores = [
            f"input[type='submit'][value='{item}']",
            f"input[type='button'][value='{item}']",
            f"input[value='{item}']",
            f"button:has-text('{item}')",
            f"a:has-text('{item}')",
        ]
        for sel in seletores:
            try:
                for el in page.query_selector_all(sel):
                    if el.is_visible():
                        logger.info(f"  ✓ Botão '{item}' encontrado. Clicando...")
                        el.click()
                        return True
            except Exception:
                pass

    for el in page.query_selector_all("button, input[type='submit'], input[type='button'], a"):
        try:
            if not el.is_visible():
                continue
            texto = (el.inner_text() or el.get_attribute("value") or "").strip().upper()
            for alvo in seletores_ou_textos:
                if alvo.upper() == texto:
                    logger.info(f"  ✓ Botão encontrado via fallback de texto: '{texto}'")
                    el.click()
                    return True
        except Exception:
            continue
    return False


def preencher_senha_e_gravar(page, senha: str) -> bool:
    """
    Preenche a senha e confirma a gravação no Sigeduc.
    Suporta:
    1. Fluxo direto: campo de senha já visível na página principal -> preenche e clica em Gravar.
    2. Fluxo com popup/modal (ex: tela de Notas): clica em Gravar -> aguarda o modal 'Confirme sua senha'
       -> preenche o input de senha no popup -> clica no botão 'Confirmar' (ou envia Enter).
    """
    if not senha:
        logger.warning("  Senha não configurada. Gravação automática cancelada.")
        return False

    logger.info("  Iniciando processo de gravação...")

    # 1. Verifica se há campo de senha visível na tela principal antes de clicar em Gravar
    campos_senha_iniciais = [
        el for el in page.query_selector_all('input[type="password"]')
        if el.is_visible()
    ]

    if campos_senha_iniciais:
        logger.info("  ✓ Campo de senha encontrado na tela principal. Preenchendo...")
        campos_senha_iniciais[0].fill(senha)
        time.sleep(0.5)
        logger.info("  Clicando em 'Gravar'...")
        if not _clicar_primeiro_visivel(page, ["Gravar", "gravar", "GRAVAR", "Salvar", "salvar", "Cadastrar"]):
            logger.error("  ✗ Botão de Gravar não encontrado após preencher senha.")
            return False
    else:
        # Fluxo com popup: o campo de senha não está na tela principal.
        # Clica em 'Gravar' para fazer o popup modal abrir.
        logger.info("  Campo de senha não está na tela principal. Clicando em 'Gravar' para abrir o popup de confirmação...")
        if not _clicar_primeiro_visivel(page, ["Gravar", "gravar", "GRAVAR", "Salvar", "salvar", "Cadastrar"]):
            logger.error("  ✗ Botão de Gravar não encontrado na tela!")
            return False

    # 2. Aguarda o popup modal 'Confirme sua senha' aparecer
    logger.info("  Aguardando popup modal de confirmação de senha...")
    time.sleep(1.0)

    campo_modal = None
    for _ in range(12):  # tenta por até 6 segundos
        for el in page.query_selector_all('input[type="password"]'):
            try:
                if el.is_visible():
                    campo_modal = el
                    break
            except Exception:
                pass
        if campo_modal:
            break
        time.sleep(0.5)

    if campo_modal:
        logger.info("  ✓ Popup 'Confirme sua senha' detectado! Digitando senha...")
        campo_modal.fill(senha)
        time.sleep(0.5)

        logger.info("  Clicando no botão 'Confirmar' do popup...")
        if _clicar_primeiro_visivel(page, ["Confirmar", "CONFIRMAR", "confirmar", "Gravar", "OK"]):
            logger.info("  ✓ Botão 'Confirmar' acionado no popup com sucesso.")
            return True
        else:
            logger.info("  Botão 'Confirmar' não clicado por seletor. Tentando enviar tecla 'Enter' no campo de senha...")
            try:
                campo_modal.press("Enter")
                logger.info("  ✓ Tecla 'Enter' enviada com sucesso no popup.")
                return True
            except Exception as e:
                logger.error(f"  ✗ Falha ao confirmar popup com Enter: {e}")
                return False
    else:
        logger.info("  Nenhum popup de senha adicional exibido (a gravação foi direta).")
        return True


# =============================================
# Processamento de Notas em Branco e Falta Vinculada (FV)
# =============================================
def processar_pergunta_notas_em_branco(page, deve_preencher_fv: bool = True, timeout: float = 6.0) -> bool:
    """
    Trata a pergunta exibida após salvar as notas sobre o que fazer com as notas em branco.
    Clica em 'Sim' se deve_preencher_fv for True, ou em 'Não' caso contrário.
    """
    logger.info("  Verificando se há pergunta sobre notas em branco pós-gravação...")
    inicio = time.time()
    alvos = ["Sim", "SIM", "sim"] if deve_preencher_fv else ["Não", "Nao", "NÃO", "nao"]

    while time.time() - inicio < timeout:
        conteudo = obter_conteudo_seguro(page).lower()
        if "falta vinculada" in conteudo or "faltas vinculadas" in conteudo:
            logger.info("  ✓ Já na tela de Falta Vinculada!")
            return True

        if _clicar_primeiro_visivel(page, alvos):
            logger.info(f"  ✓ Botão '{alvos[0]}' clicado na pergunta de notas em branco.")
            time.sleep(1.0)
            return True

        time.sleep(0.5)

    logger.info("  Nenhuma pergunta de notas em branco detectada no tempo limite.")
    return False


def preencher_faltas_vinculadas(page, dados_notas: dict, senha: str = "", unidade: str | int = None) -> int:
    """
    Preenche a tela de Falta Vinculada (FV) no SIGEduc.
    Garante a aba correta da unidade ('1ª Unidade', '2ª Unidade', etc.) antes de marcar,
    marca as checkboxes de FV para as atividades em branco dos alunos ausentes,
    digita a senha no popup de confirmação de FV e clica em Confirmar.
    Retorna a quantidade de marcações de FV realizadas.
    """
    logger.info("  Aguardando tela de atribuição de Falta Vinculada (FV)...")
    try:
        page.wait_for_load_state("networkidle", timeout=6000)
    except Exception:
        pass

    conteudo = obter_conteudo_seguro(page).lower()
    if not any(t in conteudo for t in ["falta vinculada", "faltas vinculadas", "atribuir", "fv", "verificar avaliações", "avaliações não realizadas"]):
        logger.info("  Tela de Falta Vinculada não detectada (gravação seguiu direta).")
        return 0

    logger.info("  ✓ Tela de Falta Vinculada detectada!")

    # Se uma unidade específica foi indicada, comuta para ela na tela de FV!
    if unidade:
        garantir_unidade_ativa(page, unidade)
        time.sleep(1.0)

    logger.info("  Analisando alunos para marcação de Falta Vinculada...")
    time.sleep(1.0)

    linhas = page.query_selector_all("table tbody tr")
    marcados = 0

    alunos_ja_processados = set()
    candidatos_nomes = list(dados_notas.keys())

    for linha in linhas:
        try:
            if not linha.is_visible():
                continue
            texto_linha = linha.inner_text().strip().upper()
        except Exception:
            continue

        if not texto_linha:
            continue

        # Localização robusta do aluno pelo nome do link/célula ou linha inteira
        nome_tela = extrair_nome_da_linha(linha)
        aluno_encontrado = encontrar_melhor_aluno(nome_tela, candidatos_nomes)
        if not aluno_encontrado and texto_linha and texto_linha != nome_tela:
            aluno_encontrado = encontrar_melhor_aluno(texto_linha, candidatos_nomes)

        if not aluno_encontrado:
            continue

        if aluno_encontrado in alunos_ja_processados:
            continue
        alunos_ja_processados.add(aluno_encontrado)

        item_aluno = dados_notas[aluno_encontrado]
        flags_fv = getattr(item_aluno, "fv", []) if hasattr(item_aluno, "fv") else item_aluno.get("fv", [])

        checkboxes = [chk for chk in linha.query_selector_all("input[type='checkbox']") if chk.is_visible()]
        if not checkboxes:
            continue

        for idx, chk in enumerate(checkboxes):
            precisa_marcar = (idx < len(flags_fv) and flags_fv[idx]) or (len(flags_fv) == 0)
            try:
                if precisa_marcar:
                    if not chk.is_checked():
                        chk.check()
                        logger.info(f"    ✓ {aluno_encontrado}: FV marcado na atividade {idx+1}")
                        marcados += 1
                else:
                    if chk.is_checked():
                        chk.uncheck()
                        logger.info(f"    - {aluno_encontrado}: FV desmarcado na atividade {idx+1}")
            except Exception as e:
                logger.warning(f"    ⚠ Erro ao alterar checkbox de FV para {aluno_encontrado}: {e}")

    logger.info(f"  Total de marcações de FV aplicadas: {marcados}")

    # Salva a tela de Falta Vinculada
    logger.info("  Confirmando gravação das Faltas Vinculadas...")
    time.sleep(0.5)

    if senha:
        logger.info("  Gravando e confirmando senha para as Faltas Vinculadas...")
        if preencher_senha_e_gravar(page, senha):
            logger.info("  ✓ Faltas Vinculadas gravadas e confirmadas com senha!")
            aguardar_pos_gravacao(page)
        else:
            logger.warning("  ⚠ Falha ao confirmar senha na gravação de FV.")
    else:
        if _clicar_primeiro_visivel(page, ["Gravar", "GRAVAR", "Salvar", "Confirmar", "Finalizar"]):
            logger.info("  ✓ Botão de Gravar da tela de FV clicado (aguardando inserção manual da senha).")
        else:
            logger.warning("  ⚠ Botão de gravação da tela de FV não encontrado.")

    return marcados

