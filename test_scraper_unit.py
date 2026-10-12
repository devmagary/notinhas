"""
Testes unitários e de integração para o módulo sigeduc_scraper.py
Verifica:
1. Extração de turnos e prevenção de falsos positivos (ex: MAT em MATEMATICA).
2. Formatação dos nomes das abas com Turma, Turno, Disciplina e Unidade (limite 31 caracteres).
3. Geração da planilha Excel com múltiplas turmas distintas e coluna de Turno no resumo.
"""

import os
import unittest
import pandas as pd
import openpyxl
import sigeduc_scraper as ss


class TestSigeducScraper(unittest.TestCase):

    def test_extrair_turno(self):
        self.assertEqual(ss.extrair_turno("MATEMÁTICA"), "")
        self.assertEqual(ss.extrair_turno("INFORMÁTICA"), "")
        self.assertEqual(ss.extrair_turno("3º ANO A - MATUTINO"), "Mat")
        self.assertEqual(ss.extrair_turno("2º ANO B (VESPERTINO)"), "Vesp")
        self.assertEqual(ss.extrair_turno("1º ANO C - NOTURNO"), "Not")
        self.assertEqual(ss.extrair_turno("TURMA INTEGRAL"), "Int")
        self.assertEqual(ss.extrair_turno("1º B VESP"), "Vesp")
        self.assertEqual(ss.extrair_turno("2º A MAT"), "Mat")
        # Códigos alfanuméricos comuns do SIGEduc BA
        self.assertEqual(ss.extrair_turno("ADMVES1SA"), "Vesp")
        self.assertEqual(ss.extrair_turno("SEGUMAT1SA"), "Mat")
        self.assertEqual(ss.extrair_turno("AGRPMAT1SA"), "Mat")
        self.assertEqual(ss.extrair_turno("ANACNOTM1A2"), "Not")

    def test_formatar_nome_aba(self):
        abas = set()
        aba1 = ss.formatar_nome_aba("3º ANO A - MATUTINO", "HISTÓRIA", "U1", abas)
        self.assertEqual(aba1, "3A_Mat_Hist_U1")
        self.assertLessEqual(len(aba1), 31)

        aba2 = ss.formatar_nome_aba("2º ANO B", "MATEMÁTICA", "U2", abas, turno="Vespertino")
        self.assertEqual(aba2, "2B_Vesp_Mat_U2")
        self.assertLessEqual(len(aba2), 31)

        aba3 = ss.formatar_nome_aba("1º ANO EMI - MATUTINO", "LÍNGUA PORTUGUESA", "U1", abas)
        self.assertEqual(aba3, "1EMI_Mat_Port_U1")
        self.assertLessEqual(len(aba3), 31)

        aba4 = ss.formatar_nome_aba("INFO 2", "INFORMÁTICA", "Freq", abas, turno="Vesp")
        self.assertEqual(aba4, "Info2_Vesp_Info_Freq")
        self.assertLessEqual(len(aba4), 31)

        # Turmas com códigos alfanuméricos técnicos:
        aba_tec1 = ss.formatar_nome_aba("SEGUMAT1SA", "EDUCAÇÃO FÍSICA", "U1", abas)
        self.assertEqual(aba_tec1, "Segu1SA_Mat_EdFis_U1")
        self.assertLessEqual(len(aba_tec1), 31)

        aba_tec2 = ss.formatar_nome_aba("ADMVES1SA", "MATEMÁTICA", "U2", abas)
        self.assertEqual(aba_tec2, "Adm1SA_Vesp_Mat_U2")
        self.assertLessEqual(len(aba_tec2), 31)

        # Teste de caracteres proibidos [ ] : ? * / \
        aba_proibidos = ss.formatar_nome_aba("3º/A [Mat]:?", "História*Teste", "U1", abas)
        for char_proibido in [":", "\\", "/", "?", "*", "[", "]"]:
            self.assertNotIn(char_proibido, aba_proibidos)

        # Teste de colisão de nomes: como 3A_Mat_Hist_U1 e 3A_Mat_Hist_U1_1 já estão em 'abas', a próxima é _2
        aba_duplicada = ss.formatar_nome_aba("3º ANO A - MATUTINO", "HISTÓRIA", "U1", abas)
        self.assertEqual(aba_duplicada, "3A_Mat_Hist_U1_2")

    def test_exportar_excel_multiplas_turmas(self):
        caminho_teste = "teste_excel_multiplas_turmas.xlsx"
        if os.path.exists(caminho_teste):
            os.remove(caminho_teste)

        dados_mock = [
            {
                "turma_info": {
                    "id": 1,
                    "ano": "2026",
                    "escola": "Colégio Estadual Modelo",
                    "oferta_ensino": "Ensino Médio",
                    "ano_serie": "3ª Série",
                    "turno": "Mat",
                    "periodicidade": "ANUAL",
                    "nome_turma": "3º ANO A - MATUTINO",
                    "componente_curricular": "HISTÓRIA",
                    "qtd_alunos": 2
                },
                "notas_unidades": {
                    1: pd.DataFrame([
                        [1, "202601", "ALUNO TURMA 1 A", "8.0", "7.5", "8.0"],
                        [2, "202602", "ALUNO TURMA 1 B", "9.0", "8.5", "9.0"],
                    ], columns=["Nº", "Matrícula", "Nome do Estudante", "Ativ 1", "Simulado", "Média"]),
                    2: pd.DataFrame([
                        [1, "202601", "ALUNO TURMA 1 A", "7.0", "8.0", "7.5"],
                        [2, "202602", "ALUNO TURMA 1 B", "8.0", "8.5", "8.2"],
                    ], columns=["Nº", "Matrícula", "Nome do Estudante", "Ativ 1", "Simulado", "Média"]),
                },
                "frequencia_df": None
            },
            {
                "turma_info": {
                    "id": 2,
                    "ano": "2026",
                    "escola": "Colégio Estadual Modelo",
                    "oferta_ensino": "Ensino Médio",
                    "ano_serie": "2ª Série",
                    "turno": "Vesp",
                    "periodicidade": "ANUAL",
                    "nome_turma": "2º ANO B - VESPERTINO",
                    "componente_curricular": "MATEMÁTICA",
                    "qtd_alunos": 2
                },
                "notas_unidades": {
                    1: pd.DataFrame([
                        [1, "202603", "ALUNO TURMA 2 C", "6.0", "7.0", "6.5"],
                        [2, "202604", "ALUNO TURMA 2 D", "5.0", "6.0", "5.5"],
                    ], columns=["Nº", "Matrícula", "Nome do Estudante", "Ativ 1", "Simulado", "Média"]),
                },
                "frequencia_df": None
            }
        ]

        scraper = ss.SigeducScraper(headless=True)
        scraper.exportar_para_excel(dados_mock, caminho_teste)

        self.assertTrue(os.path.exists(caminho_teste))

        wb = openpyxl.load_workbook(caminho_teste)
        sheet_names = wb.sheetnames

        # Verifica aba de resumo
        self.assertIn("Resumo_Turmas", sheet_names)
        ws_resumo = wb["Resumo_Turmas"]
        # Verifica se o cabeçalho tem 'Turno'
        headers = [ws_resumo.cell(row=2, column=c).value for c in range(1, 11)]
        self.assertIn("Turno", headers)

        # Verifica que as turmas têm abas distintas e nomes adequados
        self.assertIn("3A_Mat_Hist_U1", sheet_names)
        self.assertIn("3A_Mat_Hist_U2", sheet_names)
        self.assertIn("2B_Vesp_Mat_U1", sheet_names)

        # Verifica que os alunos da Turma 1 não são os mesmos da Turma 2
        ws_t1 = wb["3A_Mat_Hist_U1"]
        ws_t2 = wb["2B_Vesp_Mat_U1"]

        aluno_t1 = ws_t1.cell(row=3, column=3).value
        aluno_t2 = ws_t2.cell(row=3, column=3).value

        self.assertEqual(aluno_t1, "ALUNO TURMA 1 A")
        self.assertEqual(aluno_t2, "ALUNO TURMA 2 C")
        self.assertNotEqual(aluno_t1, aluno_t2)

        wb.close()
        if os.path.exists(caminho_teste):
            os.remove(caminho_teste)

    def test_esta_no_painel_docente_criterios(self):
        """Valida que a tela de notas NUNCA é confundida com o painel de turmas."""
        class MockElement:
            def __init__(self, text="", is_checked_val=False):
                self._text = text
                self._is_checked = is_checked_val
            def inner_text(self):
                return self._text
            def get_attribute(self, attr):
                return ""
            def is_visible(self):
                return True

        class MockPageNotes:
            url = "https://sigeduc.educacao.ba.gov.br/sigeduc/portais/docente/docente.jsf"
            def content(self):
                return "<html><body>Diário de Classe Digital > LANÇAR NOTAS PARA AVALIAÇÕES DAS UNIDADES</body></html>"
            def query_selector(self, sel):
                # Simula presença do campo de nota ou tabela de resultado
                if "av__" in sel or "result" in sel:
                    return MockElement()
                return None
            def query_selector_all(self, sel):
                if sel == "table":
                    return [MockElement("Escola: Colégio Modelo Turma: ADMVES1SA Componente: HISTORIA 1ª SÉRIE")]
                return []

        class MockPageClassList:
            url = "https://sigeduc.educacao.ba.gov.br/sigeduc/portais/docente/docente.jsf"
            def content(self):
                return "<html><body>Minhas Turmas: Turmas Regulares</body></html>"
            def query_selector(self, sel):
                if "notas_lancar" in sel or "frequencia_lancar" in sel:
                    return MockElement()
                return None
            def query_selector_all(self, sel):
                if sel == "table":
                    return [MockElement("Ano Oferta Turma Componente Curricular Qtd. de Estudantes Ações")]
                return []

        scraper = ss.SigeducScraper(headless=True)
        
        # Teste 1: Na tela de notas, deve retornar estritamente FALSE
        scraper.page = MockPageNotes()
        self.assertFalse(scraper.esta_no_painel_docente(), "Tela de notas NÃO pode ser considerada painel docente!")

        # Teste 2: Na tela de listagem de turmas, deve retornar TRUE
        scraper.page = MockPageClassList()
        self.assertTrue(scraper.esta_no_painel_docente(), "Lista de turmas deve ser reconhecida com sucesso!")


if __name__ == "__main__":
    unittest.main()
